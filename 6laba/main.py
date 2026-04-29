import tkinter as tk
from tkinter import messagebox, ttk
from calculations.models import Room, Apartment, Building
from calculations.area import square
from calculations.heat import power_object
from calculations.db import db

def show_history():
    rows = db.get_all()
    if not rows:
        messagebox.showinfo("История", "Нет расчетов")
        return
    
    win = tk.Toplevel(root)
    win.title("История")
    win.geometry("500x300")
    
    text = tk.Text(win)
    text.pack(fill="both", expand=True)
    
    for row in rows:
        text.insert("end", f"{row[0]}. {row[1]} | {row[2]} | {row[3]}м² | {row[4]}Вт | {row[5]}\n")
        text.insert("end", "-" * 40 + "\n")
    
    text.config(state="disabled")

def calculate():
    global last_result
    
    if var.get() == 1:
        w = float(entry_w.get())
        l = float(entry_l.get())
        obj = Room(w, l)
        area = square(obj)
        heat = power_object(obj)
        result = f"Комната: {w}x{l} = {area:.2f}м², {heat:.2f}Вт"
        db.save("комната", f"{w}x{l}", area, heat)
    
    elif var.get() == 2:
        data = text_rooms.get("1.0", tk.END).strip()
        obj = Apartment()
        rooms_list = []
        for line in data.split('\n'):
            if line.strip():
                w, l = map(float, line.split())
                obj.rooms.append(Room(w, l))
                rooms_list.append(f"{w}x{l}")
        area = square(obj)
        heat = power_object(obj)
        result = f"Квартира, {len(obj.rooms)} комнат, площадь {area:.2f}м², тепло {heat:.2f}Вт"
        db.save("квартира", " ".join(rooms_list), area, heat)
    
    else:
        floors = int(entry_floors.get())
        count = int(entry_count.get())
        avg = float(entry_avg.get())
        obj = Building()
        side = avg ** 0.5
        for i in range(count):
            apt = Apartment()
            apt.rooms.append(Room(side, side))
            obj.apartments.append(apt)
        area = square(obj)
        heat = power_object(obj)
        result = f"Дом: {floors} эт, {count} кв, площадь {area:.2f}м², тепло {heat:.2f}Вт"
        db.save("дом", f"{floors}эт {count}кв", area, heat)
    
    messagebox.showinfo("Сохранено", "Результат сохранен в БД")
    label_result.config(text=result)
    last_result = result

def save_to_file():
    if not last_result:
        messagebox.showwarning("Ошибка", "Сначала рассчитайте")
        return
    
    win = tk.Toplevel(root)
    win.title("Сохранить")
    win.geometry("200x100")
    
    tk.Button(win, text="Word", command=lambda: [save_word(last_result), win.destroy()]).pack(pady=5)
    tk.Button(win, text="Excel", command=lambda: [save_excel(last_result), win.destroy()]).pack(pady=5)

def save_word(text):
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    doc.save("report.docx")
    messagebox.showinfo("Успех", "Сохранено в report.docx")

def save_excel(text):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A1"] = text
    wb.save("report.xlsx")
    messagebox.showinfo("Успех", "Сохранено в report.xlsx")

root = tk.Tk()
root.title("Расчет помещений")
root.geometry("500x450")
root.resizable(False, False)

last_result = ""
var = tk.IntVar(value=1)

tk.Radiobutton(root, text="Комната", variable=var, value=1).pack()
tk.Radiobutton(root, text="Квартира", variable=var, value=2).pack()
tk.Radiobutton(root, text="Дом", variable=var, value=3).pack()

frame = tk.Frame(root)
frame.pack()

entry_w = tk.Entry(frame, width=10)
entry_l = tk.Entry(frame, width=10)
tk.Label(frame, text="Ширина:").grid(row=0, column=0)
entry_w.grid(row=0, column=1)
tk.Label(frame, text="Длина:").grid(row=1, column=0)
entry_l.grid(row=1, column=1)

text_rooms = tk.Text(frame, height=5, width=30)
tk.Label(frame, text="Комнаты (ширина длина):").grid(row=2, column=0, columnspan=2)
text_rooms.grid(row=3, column=0, columnspan=2)

entry_floors = tk.Entry(frame, width=10)
entry_count = tk.Entry(frame, width=10)
entry_avg = tk.Entry(frame, width=10)
tk.Label(frame, text="Этажи:").grid(row=4, column=0)
entry_floors.grid(row=4, column=1)
tk.Label(frame, text="Квартир:").grid(row=5, column=0)
entry_count.grid(row=5, column=1)
tk.Label(frame, text="Ср.площадь:").grid(row=6, column=0)
entry_avg.grid(row=6, column=1)

def switch():
    if var.get() == 1:
        entry_w.grid()
        entry_l.grid()
        text_rooms.grid_remove()
        entry_floors.grid_remove()
        entry_count.grid_remove()
        entry_avg.grid_remove()
    elif var.get() == 2:
        entry_w.grid_remove()
        entry_l.grid_remove()
        text_rooms.grid()
        entry_floors.grid_remove()
        entry_count.grid_remove()
        entry_avg.grid_remove()
    else:
        entry_w.grid_remove()
        entry_l.grid_remove()
        text_rooms.grid_remove()
        entry_floors.grid()
        entry_count.grid()
        entry_avg.grid()

var.trace("w", lambda *args: switch())
switch()

tk.Button(root, text="Рассчитать", command=calculate).pack(pady=5)
tk.Button(root, text="Сохранить в файл", command=save_to_file).pack(pady=5)
tk.Button(root, text="История БД", command=show_history).pack(pady=5)
label_result = tk.Label(root, text="", justify=tk.LEFT)
label_result.pack(pady=10)

def on_closing():
    db.close()
    root.destroy()

root.protocol("WM_DELETE_WINDOW", on_closing)
root.mainloop()